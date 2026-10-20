import openpyxl

wb = openpyxl.load_workbook("daily_reports.xlsx", data_only=True)

# Берём один точно рабочий лист из списка
target_sheet = "24.10.24 NOON AT SEA"
if target_sheet not in wb.sheetnames:
    target_sheet = [s for s in wb.sheetnames if "NOON AT SEA" in s][0]

sheet = wb[target_sheet]
print(f"--- Диагностика листа: {sheet.title} ---")

for row in sheet.iter_rows(values_only=False):
    for cell in row:
        val = str(cell.value).strip() if cell.value is not None else ""
        if "MAIN ENGINE" in val.upper() or "CONSUMPTION" in val.upper() or "VLSFO" in val.upper():
            print(f"Найдено '{val}' в строке {cell.row}, колонке {cell.column}")
            # Печатаем все непустые ячейки в этой строке
            row_vals = [f"Col {c.column}: {c.value}" for c in sheet[cell.row] if c.value is not None]
            print("  Вся строка:", " | ".join(row_vals))